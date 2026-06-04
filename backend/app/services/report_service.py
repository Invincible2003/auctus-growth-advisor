import os
import logging
from datetime import datetime
from app.core.config import settings

# Attempt imports for PDF, DOCX, PPTX
try:
    from fpdf import FPDF
    has_fpdf = True
except ImportError:
    has_fpdf = False

try:
    import docx
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    has_docx = True
except ImportError:
    has_docx = False

try:
    import pptx
    from pptx.util import Inches as PptInches, Pt as PptPt
    from pptx.dml.color import RGBColor
    has_pptx = True
except ImportError:
    has_pptx = False

logger = logging.getLogger(__name__)

# Output directory for reports
REPORTS_DIR = "D:/AUCTUS/reports"
if not os.path.exists(REPORTS_DIR):
    os.makedirs(REPORTS_DIR, exist_ok=True)

class BrandedPDF(FPDF if has_fpdf else object):
    def header(self):
        if not has_fpdf:
            return
        # Background color bar
        self.set_fill_color(11, 25, 44) # Dark blue (#0B192C)
        self.rect(0, 0, 210, 15, 'F')
        
        self.set_text_color(255, 255, 255)
        self.set_font('Arial', 'B', 10)
        self.cell(0, -6, f"{settings.PROJECT_NAME} - AI Growth Report", 0, 0, 'L')
        self.cell(0, -6, settings.BRAND_OWNER, 0, 1, 'R')
        self.ln(10)

    def footer(self):
        if not has_fpdf:
            return
        self.set_y(-15)
        self.set_text_color(128, 128, 128)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f"Page {self.page_no()} | © 2026 AUCTUS | {settings.BRAND_OWNER}", 0, 0, 'C')

class ReportService:
    @classmethod
    def generate_pdf_report(cls, business_name: str, kpis: dict, recommendations: list, filename: str) -> str:
        """Generates a branded executive PDF report."""
        file_path = os.path.join(REPORTS_DIR, filename)
        
        if not has_fpdf:
            # Fallback mock empty file creation
            with open(file_path, "w") as f:
                f.write(f"PDF generation fallback. Business: {business_name}")
            return file_path
            
        try:
            pdf = BrandedPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=20)
            
            # Title
            pdf.set_text_color(11, 25, 44)
            pdf.set_font('Arial', 'B', 22)
            pdf.cell(0, 15, "AUCTUS Growth Performance Report", 0, 1, 'C')
            pdf.set_font('Arial', 'I', 11)
            pdf.cell(0, 5, f"Prepared for: {business_name} | Date: {datetime.now().strftime('%Y-%m-%d')}", 0, 1, 'C')
            pdf.ln(10)
            
            # Executive Summary Section
            pdf.set_fill_color(240, 245, 250)
            pdf.rect(10, 45, 190, 25, 'F')
            pdf.set_xy(12, 47)
            pdf.set_font('Arial', 'B', 12)
            pdf.set_text_color(30, 41, 59)
            pdf.cell(0, 6, "Executive Summary", 0, 1)
            pdf.set_font('Arial', '', 9)
            pdf.multi_cell(0, 5, (
                f"This data-driven growth audit was automatically compiled by AUCTUS. "
                f"Based on historical transactional behavior, customer feedback logs, and competitive positioning, "
                f"we have identified multiple tactical actions to maximize monthly profit margins and acquire new local market share."
            ))
            pdf.ln(12)
            
            # Key Performance Indicators (KPIs)
            pdf.set_font('Arial', 'B', 14)
            pdf.set_text_color(11, 25, 44)
            pdf.cell(0, 8, "Business Metrics & KPIs", 0, 1)
            pdf.ln(2)
            
            # KPI Table Headers
            pdf.set_font('Arial', 'B', 10)
            pdf.set_fill_color(21, 76, 121) # Cyan-blue accent
            pdf.set_text_color(255, 255, 255)
            pdf.cell(90, 7, "Metric Description", 1, 0, 'L', True)
            pdf.cell(100, 7, "Current Value", 1, 1, 'C', True)
            
            # Table Data
            pdf.set_text_color(30, 41, 59)
            pdf.set_font('Arial', '', 10)
            for key, val in kpis.items():
                pdf.cell(90, 7, str(key), 1, 0, 'L')
                pdf.cell(100, 7, str(val), 1, 1, 'C')
            
            pdf.ln(10)
            
            # Marketing Recommendations
            pdf.set_font('Arial', 'B', 14)
            pdf.set_text_color(11, 25, 44)
            pdf.cell(0, 8, "AI growth Recommendations", 0, 1)
            pdf.ln(2)
            
            for i, rec in enumerate(recommendations, 1):
                pdf.set_font('Arial', 'B', 11)
                pdf.set_text_color(21, 76, 121)
                pdf.cell(0, 6, f"{i}. {rec.get('title', 'Recommendation')}", 0, 1)
                
                pdf.set_font('Arial', '', 9)
                pdf.set_text_color(30, 41, 59)
                pdf.multi_cell(0, 5, f"Impact: {rec.get('expected_impact', 'High')} | Difficulty: {rec.get('implementation_difficulty', 'Easy')} | Score: {rec.get('priority_score', 80)}/100")
                pdf.multi_cell(0, 4.5, f"Plan: {rec.get('description', '')}")
                pdf.ln(4)
                
            # Branded Footer Signature
            pdf.ln(5)
            pdf.set_draw_color(200, 200, 200)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(5)
            pdf.set_font('Arial', 'B', 10)
            pdf.set_text_color(11, 25, 44)
            pdf.cell(0, 6, "AUCTUS Growth Consultant Team", 0, 1, 'C')
            pdf.set_font('Arial', '', 9)
            pdf.cell(0, 5, f"Final Year Project Portfolio - {settings.BRAND_OWNER}", 0, 1, 'C')
            
            pdf.output(file_path)
            return file_path
        except Exception as e:
            logger.error(f"Failed to generate PDF report: {e}")
            return file_path

    @classmethod
    def generate_docx_report(cls, business_name: str, kpis: dict, recommendations: list, filename: str) -> str:
        """Generates a formatted Word Document report."""
        file_path = os.path.join(REPORTS_DIR, filename)
        
        if not has_docx:
            with open(file_path, "w") as f:
                f.write(f"Word generation fallback. Business: {business_name}")
            return file_path
            
        try:
            doc = docx.Document()
            
            # Header
            header = doc.sections[0].header
            hp = header.paragraphs[0]
            hp.text = f"{settings.PROJECT_NAME} Advisory Report | Brand Owner: {settings.BRAND_OWNER}"
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hp.style.font.size = Pt(8.5)
            hp.style.font.italic = True
            
            # Title
            title = doc.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title.add_run("AUCTUS Growth Performance Report")
            run.font.name = 'Arial'
            run.font.size = Pt(24)
            run.font.bold = True
            
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_run = sub.add_run(f"Prepared for: {business_name} | Date: {datetime.now().strftime('%Y-%m-%d')}")
            sub_run.font.size = Pt(11)
            sub_run.font.italic = True
            
            # Intro
            doc.add_heading("Executive Summary", level=1)
            doc.add_paragraph(
                "This performance report was generated automatically by AUCTUS. "
                "By tracking customer review scores, sales transaction rates, and competitor profiles, "
                "AUCTUS outlines recommendations for inventory planning, ad spending, and local digital marketing."
            )
            
            # KPIs Table
            doc.add_heading("Key Business KPIs", level=1)
            table = doc.add_table(rows=1, cols=2)
            table.style = 'Light Shading Accent 1'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'KPI Metric'
            hdr_cells[1].text = 'Current Rating'
            
            for key, val in kpis.items():
                row_cells = table.add_row().cells
                row_cells[0].text = str(key)
                row_cells[1].text = str(val)
                
            # Recommendations
            doc.add_heading("AI Growth Advisor Action Plan", level=1)
            for rec in recommendations:
                p = doc.add_paragraph()
                r_title = p.add_run(f"• {rec.get('title', 'Recommendation')}")
                r_title.bold = True
                p.add_run(f"\nExpected Impact: {rec.get('expected_impact')} | Ease: {rec.get('implementation_difficulty')}\n")
                p.add_run(f"Action Detail: {rec.get('description')}\n")
                
            # Footer
            doc.add_paragraph("\n\nReport generated by AUCTUS, created by Aryan Pandey.")
            
            doc.save(file_path)
            return file_path
        except Exception as e:
            logger.error(f"Failed to generate Word report: {e}")
            return file_path

    @classmethod
    def generate_pptx_report(cls, business_name: str, kpis: dict, recommendations: list, filename: str) -> str:
        """Generates a PowerPoint presentation report."""
        file_path = os.path.join(REPORTS_DIR, filename)
        
        if not has_pptx:
            with open(file_path, "w") as f:
                f.write(f"PowerPoint generation fallback. Business: {business_name}")
            return file_path
            
        try:
            prs = pptx.Presentation()
            
            # Slide 1: Title Slide
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            title.text = f"{settings.PROJECT_NAME} Advisory"
            subtitle.text = f"Growth & Sales Diagnostics\nBusiness: {business_name}\nCreated by Aryan Pandey"
            
            # Slide 2: KPI Dashboard Slide
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            title_shape.text = "Key Performance Indicators"
            
            # Add table
            rows = len(kpis) + 1
            cols = 2
            left = PptInches(1.0)
            top = PptInches(2.0)
            width = PptInches(8.0)
            height = PptInches(4.0)
            
            tbl_shape = shapes.add_table(rows, cols, left, top, width, height)
            table = tbl_shape.table
            table.columns[0].width = PptInches(4.5)
            table.columns[1].width = PptInches(3.5)
            
            table.cell(0, 0).text = 'Metric Name'
            table.cell(0, 1).text = 'Current Level'
            
            for idx, (k, v) in enumerate(kpis.items(), 1):
                table.cell(idx, 0).text = str(k)
                table.cell(idx, 1).text = str(v)
                
            # Slide 3: Growth Roadmap Slide
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            title_shape.text = "AI Recommended Roadmap"
            
            body_shape = shapes.placeholders[1]
            tf = body_shape.text_frame
            tf.text = "Identified Growth Strategies:"
            
            for rec in recommendations[:3]: # Add top 3
                p = tf.add_paragraph()
                p.text = f"• {rec.get('title')}: {rec.get('description')[:120]}..."
                p.level = 1
                
            # Add copyright/branding label to slide 3
            txBox = slide.shapes.add_textbox(PptInches(1.0), PptInches(6.5), PptInches(8.0), PptInches(0.5))
            tf_c = txBox.text_frame
            p_c = tf_c.add_paragraph()
            p_c.text = f"© 2026 AUCTUS | {settings.BRAND_OWNER}"
            p_c.font.size = PptPt(10)
            p_c.font.color.rgb = RGBColor(128, 128, 128)
            
            prs.save(file_path)
            return file_path
        except Exception as e:
            logger.error(f"Failed to generate PPTX report: {e}")
            return file_path
